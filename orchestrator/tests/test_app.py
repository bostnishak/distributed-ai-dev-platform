"""Orkestratorun harici bir model/gateway'e ihtiyac duymayan, saf mantik iceren
kisimlari icin birim testleri. Gercek Ollama/LiteLLM entegrasyonu bu testlerde
CALISTIRILMIYOR -- CI ortaminda GPU/model olmadigi icin bu testler her zaman
calisabilir olmali."""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from fastapi.testclient import TestClient

from app import (
    app,
    extract_docx_text,
    extract_pdf_text,
    keyword_category_override,
    parse_classifier_response,
)


class TestKeywordCategoryOverride:
    def test_detects_python_keyword(self):
        assert keyword_category_override("Python'da bir fonksiyon yaz") == "kod"

    def test_detects_uygulama_yap(self):
        assert keyword_category_override("Basit bir hesap makinesi uygulamasi yap") == "kod"

    def test_detects_debug_request(self):
        assert keyword_category_override("Bu kodda hata var, debug et") == "kod"

    def test_no_match_for_general_question(self):
        assert keyword_category_override("Fotosentez nedir?") is None

    def test_no_match_for_emotional_message(self):
        assert keyword_category_override("Isimden kovuldum, ne yapmaliyim?") is None

    def test_case_insensitive(self):
        assert keyword_category_override("PYTHON ile bir script yaz") == "kod"


class TestParseClassifierResponse:
    def test_clean_format(self):
        category, inappropriate = parse_classifier_response("KATEGORI: kod\nUYGUNSUZ: hayir")
        assert category == "kod"
        assert inappropriate is False

    def test_inappropriate_true(self):
        category, inappropriate = parse_classifier_response("KATEGORI: genel\nUYGUNSUZ: evet")
        assert category == "genel"
        assert inappropriate is True

    def test_handles_markdown_bold(self):
        # Gercek testte model bazen boyle kalin yazi ekliyordu (2026-09-25).
        category, inappropriate = parse_classifier_response("KATEGORI: **kod**\nUYGUNSUZ: **hayir**")
        assert category == "kod"
        assert inappropriate is False

    def test_template_echo_does_not_crash(self):
        # Gercek testte model sablonu oldugu gibi kopyalamisti (2026-09-25) -- boyle bozuk
        # bir cikti hata firlatmamali, gecerli bir kategoriyle (regex neyi bulursa) devam
        # etmeli, ve UYGUNSUZ alani net oldugu icin yanlislikla engellenmemeli.
        category, inappropriate = parse_classifier_response("KATEGORI: kod|analiz|genel\nUYGUNSUZ: hayir")
        assert category in ("kod", "genel")
        assert inappropriate is False

    def test_missing_uygunsuz_field_fails_open(self):
        category, inappropriate = parse_classifier_response("KATEGORI: kod")
        assert category == "kod"
        assert inappropriate is False

    def test_completely_malformed_fails_open(self):
        category, inappropriate = parse_classifier_response("bunun hicbir anlami yok")
        assert category == "genel"
        assert inappropriate is False

    def test_evet_appearing_elsewhere_does_not_trigger_false_positive(self):
        # Gercek testte "evet" kelimesi UYGUNSUZ satirindan uzakta gecince yanlislikla
        # engelleme tetikleniyordu -- SIKI regex bunu onlemeli.
        text = "KATEGORI: kod\nUYGUNSUZ: hayir\nEvet, bu kod ornegi calisir."
        _category, inappropriate = parse_classifier_response(text)
        assert inappropriate is False


class TestFileExtraction:
    def test_extract_docx_text(self):
        doc = Document()
        doc.add_paragraph("Test satiri bir")
        doc.add_paragraph("Test satiri iki")
        buf = io.BytesIO()
        doc.save(buf)
        text = extract_docx_text(buf.getvalue())
        assert "Test satiri bir" in text
        assert "Test satiri iki" in text

    def test_extract_pdf_text_does_not_crash_on_blank_pdf(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buf = io.BytesIO()
        writer.write(buf)
        text = extract_pdf_text(buf.getvalue())
        assert isinstance(text, str)


class TestHealthEndpoint:
    def test_health_ok(self):
        client = TestClient(app)
        resp = client.get("/health")
        assert resp.status_code == 200
        assert resp.json() == {"status": "ok"}

    def test_chat_ui_served(self):
        client = TestClient(app)
        resp = client.get("/")
        assert resp.status_code == 200
        assert "text/html" in resp.headers["content-type"]
