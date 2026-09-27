import io

import pytest
from docx import Document
from pypdf import PdfWriter

from extract import UnsupportedFileError, extract_docx_text, extract_pdf_text, extract_text


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_text():
    text = extract_docx_text(_docx_bytes(["Test satiri bir", "Test satiri iki"]))
    assert "Test satiri bir" in text
    assert "Test satiri iki" in text


def test_extract_docx_includes_tables():
    text = extract_docx_text(_docx_bytes(["Gereksinimler"], [["ID", "Açıklama"], ["FR-1", "Kitap ekleme"]]))
    assert "FR-1 | Kitap ekleme" in text


def test_extract_pdf_text_does_not_crash_on_blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    assert isinstance(extract_pdf_text(buf.getvalue()), str)


def test_extract_text_dispatches_by_extension():
    assert extract_text("notlar.MD", "# Başlık".encode()) == "# Başlık"
    assert extract_text("a.txt", b"plain") == "plain"
    assert "satir" in extract_text("a.docx", _docx_bytes(["satir"]))


def test_extract_text_rejects_unknown_types():
    with pytest.raises(UnsupportedFileError):
        extract_text("resim.png", b"\x89PNG")
