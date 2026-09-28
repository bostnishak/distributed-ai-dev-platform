"""Plain-text extraction for uploaded documents (requirements documents and chat attachments)."""

import csv
import io

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

REQUIREMENT_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class UnsupportedFileError(ValueError):
    pass


def extract_pdf_text(raw: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx_text(raw: bytes) -> str:
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    # Requirements documents often keep details in tables; paragraphs alone would lose them.
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_csv_text(raw: bytes) -> str:
    text = raw.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    # Same "|"-separated shape as a Markdown table, so it reads cleanly in a prompt and the
    # model can echo it back in the same format.
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def extract_xlsx_text(raw: bytes) -> str:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"--- Sayfa: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def extract_text(filename: str, raw: bytes) -> str:
    """Return the text of a PDF, DOCX, TXT or Markdown file; raise UnsupportedFileError otherwise."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_pdf_text(raw)
    if name.endswith(".docx"):
        return extract_docx_text(raw)
    if name.endswith((".txt", ".md")):
        return raw.decode("utf-8", errors="replace")
    raise UnsupportedFileError(filename)
